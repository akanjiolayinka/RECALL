"""Generate Recall's synthetic demo/test files into test-data/.

Every document here is invented. Never put real personal files in test-data/.

Usage (needs the dev-only packages reportlab, python-docx and Pillow):
    python -m venv .venv-testdata
    .venv-testdata/bin/pip install reportlab python-docx pillow   # Windows: .venv-testdata\\Scripts\\pip
    .venv-testdata/bin/python scripts/generate_test_data.py

The output is committed, so most contributors never need to run this.
"""

from datetime import datetime
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent.parent / "test-data"
FIXED_DATE = datetime(2026, 5, 2, 9, 0, 0)


def write_pdf(path: Path, title: str, author: str, pages: list[list[str]]) -> None:
    """Write a simple text PDF, one list of lines per page."""
    path.parent.mkdir(parents=True, exist_ok=True)
    # invariant=1 makes reportlab's output byte-for-byte reproducible.
    pdf = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    pdf.setTitle(title)
    pdf.setAuthor(author)
    width, height = A4
    for number, lines in enumerate(pages, start=1):
        y = height - 72
        for line in lines:
            if line.startswith("# "):
                pdf.setFont("Helvetica-Bold", 16)
                line = line[2:]
            else:
                pdf.setFont("Helvetica", 11)
            pdf.drawString(72, y, line)
            y -= 22 if line else 12
        pdf.setFont("Helvetica", 9)
        pdf.drawString(width / 2 - 20, 40, f"Page {number} of {len(pages)}")
        pdf.showPage()
    pdf.save()


def write_image_only_pdf(path: Path) -> None:
    """A PDF with shapes but no text layer, like a scanned page."""
    pdf = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    pdf.rect(72, 500, 450, 250)
    pdf.line(90, 700, 480, 700)
    pdf.line(90, 660, 400, 660)
    pdf.showPage()
    pdf.save()


def write_docx(path: Path, title: str, author: str, paragraphs: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.author = author
    doc.core_properties.created = FIXED_DATE
    doc.core_properties.modified = FIXED_DATE
    for paragraph in paragraphs:
        if paragraph.startswith("# "):
            doc.add_heading(paragraph[2:], level=1)
        else:
            doc.add_paragraph(paragraph)
    doc.save(str(path))


def write_receipt_image(path: Path, lines: list[str]) -> None:
    """A plain receipt-like PNG: dark text on white, for OCR testing."""
    path.parent.mkdir(parents=True, exist_ok=True)
    font = ImageFont.load_default(size=28)
    image = Image.new("RGB", (720, 80 + 46 * len(lines)), "white")
    draw = ImageDraw.Draw(image)
    for i, line in enumerate(lines):
        draw.text((40, 40 + 46 * i), line, fill="black", font=font)
    image.save(path, optimize=False)


def write_text(path: Path, text: str, encoding: str = "utf-8") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(text.encode(encoding))


def main() -> None:
    write_pdf(
        ROOT / "pdf" / "Project Proposal - Riverside Community Garden.pdf",
        "Riverside Community Garden - Project Proposal",
        "Amaka Obi",
        [
            [
                "# Riverside Community Garden",
                "Project proposal prepared for the Riverside Residents Association.",
                "",
                "Summary",
                "We propose turning the empty lot behind Block C into a shared vegetable",
                "garden with raised beds, a rainwater tank and a small tool shed.",
            ],
            [
                "# Goals and timeline",
                "Phase one (March): clear the lot and test the soil.",
                "Phase two (April): build twelve raised beds and install the water tank.",
                "Phase three (May): open the garden to members and run a planting day.",
            ],
            [
                "# Budget",
                "The estimated project budget is NGN 2,500,000 in total.",
                "Raised beds and soil: NGN 1,100,000",
                "Rainwater tank and pipes: NGN 650,000",
                "Tool shed and tools: NGN 450,000",
                "Contingency (about 12 percent): NGN 300,000",
                "Funding will come from member dues and a ward development grant.",
            ],
            [
                "# Risks",
                "The main risks are a dry season water shortage and theft of tools.",
                "We will lock the shed and keep a sign-out sheet for equipment.",
            ],
        ],
    )
    write_pdf(
        ROOT / "pdf" / "Tenancy Agreement - Flat 4B.pdf",
        "Tenancy Agreement - Flat 4B",
        "Lekki Homes Ltd",
        [
            [
                "# Tenancy Agreement",
                "Property: Flat 4B, 12 Palm Avenue.",
                "Tenant: Tunde Bello. Landlord: Lekki Homes Ltd.",
                "Term: twelve months starting 1 June 2026.",
            ],
            [
                "# Deposit and notice",
                "The tenant pays a refundable security deposit of two months' rent.",
                "Either party must give sixty days' written notice before moving out.",
                "The deposit is returned within thirty days after the keys are handed back,",
                "minus the cost of repairing any damage beyond normal wear and tear.",
            ],
        ],
    )
    write_image_only_pdf(ROOT / "pdf" / "Scanned letter (no text layer).pdf")

    write_docx(
        ROOT / "documents" / "Q3 Marketing Plan.docx",
        "Q3 Marketing Plan",
        "Chidi Eze",
        [
            "# Q3 Marketing Plan",
            "This quarter we focus on the back-to-school campaign and the new mobile app launch.",
            "# Spending",
            "Total marketing spend for the quarter is capped at NGN 4,000,000, "
            "split between radio adverts, social media and two community events.",
            "# Success measures",
            "We will track app downloads, newsletter sign-ups and event attendance every week.",
        ],
    )

    write_text(
        ROOT / "notes" / "Moving out checklist.md",
        "# Moving out checklist\n\n"
        "- Give the landlord sixty days' notice in writing\n"
        "- Book the removal van for the last Saturday of the month\n"
        "- Take photos of every room before handing back the keys\n"
        "- Ask when the security deposit will be refunded\n"
        "- Redirect post to the new address\n",
    )
    write_text(
        ROOT / "notes" / "Meeting notes 2026-05-02.txt",
        "Garden committee meeting, 2 May 2026\n\n"
        "Present: Amaka, Tunde, Ngozi.\n\n"
        "We agreed to order the rainwater tank next week. Ngozi will ask the ward office "
        "about the development grant. The planting day is pencilled in for 23 May.\n",
    )
    write_text(
        ROOT / "notes" / "Shopping list (UTF-16).txt",
        "Shopping list: tomatoes, peppers, garden gloves, watering can.\n",
        encoding="utf-16",  # writes a byte-order mark, like Windows Notepad's "Unicode"
    )
    write_receipt_image(
        ROOT / "images" / "Headphones receipt.png",
        [
            "SOUNDWAVE ELECTRONICS",
            "14 Market Road, Ikeja",
            "Date: 12/04/2026",
            "Wireless headphones   NGN 45,000",
            "Carrying case         NGN  5,000",
            "TOTAL                 NGN 50,000",
            "Paid by card. Thank you!",
        ],
    )
    print(f"Test data written to {ROOT}")


if __name__ == "__main__":
    main()
